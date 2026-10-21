"""Build a Thai DOCX using fresh outputs; no fabricated or historical p-values."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
OUT=ROOT/"outputs"
def create():
    c=json.loads((OUT/"classification.json").read_text())
    r=json.loads((OUT/"regression.json").read_text())
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.27);s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.65)
    s.left_margin=s.right_margin=Inches(.7)
    for name in ["Normal","Title","Heading 1"]:
        st=doc.styles[name];st.font.name="Garuda";st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(10.5 if name=="Normal" else 14 if name=="Title" else 11.5)
        st.paragraph_format.space_after=Pt(5);st.paragraph_format.space_before=Pt(0)
        st.paragraph_format.line_spacing=1.05
        rp=st.element.get_or_add_rPr()
        rf=rp.find(qn("w:rFonts"))
        if rf is None: rf=OxmlElement("w:rFonts");rp.append(rf)
        for key in ["ascii","hAnsi","eastAsia","cs"]:rf.set(qn("w:"+key),"Garuda")
        for key in ["asciiTheme","hAnsiTheme","eastAsiaTheme","cstheme"]:
            if qn("w:"+key) in rf.attrib: del rf.attrib[qn("w:"+key)]
    def p(text):doc.add_paragraph(text)
    def h(text):doc.add_heading(text,1)
    doc.add_heading("การพัฒนา FFN สำหรับข้อมูล SES2023",0)
    h("1 Introduction")
    p("งานนี้ทำสำเนาโมเดลตัวอย่าง classification และ regression เพื่อพัฒนา Feedforward Neural Networks สำหรับทำนายสถานะความยากจนและค่าใช้จ่ายบริโภคต่อหัว ข้อมูล SES2023 มี 49,664 รายการ กลุ่มยากจนมี 769 รายการ หรือ 1.55% จึงประเมิน classification ด้วย Precision, Recall, F1 และ PR-AUC ร่วมกับ Accuracy ส่วน regression ใช้ RMSE, MAE และ R²")
    h("2 Model Architecture and Conceptual Framework")
    p("คงโครงสร้าง 32–64–32–1 และ ReLU การพัฒนาเน้นวิธีเรียนรู้ ไม่ใช่การเพิ่มจำนวนชั้น Classification ใช้ Weighted BCE โดยเพิ่ม pos_weight จาก 1 เป็น 5 ส่วน regression เปลี่ยน MSE เป็น SmoothL1Loss แบบ Huber โดยกำหนด beta=1 หลัง standardize target ทั้งสองงานใช้ Adam, learning rate 0.001 และเพิ่ม batch size เป็น 512 จากเดิม 64 และ 32 ตามลำดับ")
    p("ใช้ Early stopping เลือก checkpoint ด้วย validation PR-AUC สำหรับ classification และ RMSE สำหรับ regression โดย patience=12 และสูงสุด 50 epochs เฉลี่ยผลจากสาม FFNs ที่ใช้ seeds 0, 1, 2 เพื่อลดความผันผวน Classification เลือก threshold จาก F1 บนชุด calibration โมเดลที่เลือกไม่เพิ่ม Dropout หรือ Batch Normalization การกำหนดค่ามาจากการศึกษาข้อมูลพัฒนารอบก่อน ไม่ได้เลือกจากผล test ของการรันนี้")
    h("3 Methods and Data")
    p("ตัด pov_exp และ rspcexp23 ออกจาก predictors ทั้งสองงาน แทน missing predictors ด้วยมัธยฐานและ standardize โดย fit เฉพาะ training rows คง test split เดิม 20% ด้วย random_state=1 ส่วนข้อมูลพัฒนาแบ่งประมาณ 60% สำหรับฝึก 10% สำหรับเลือก epoch/threshold และ 10% สำหรับเปรียบเทียบวิธี โดย stratify ตามความยากจน ชุด 10% หลังถูกสงวนไว้ในการรัน configuration ที่เลือกแล้ว")
    p("รัน baseline ที่รวมในไฟล์แต่ละงาน โดยคงวิธีฝึกของ LAB เดิมและใช้ CPU ทั้งสองแบบ Baseline ยังคงใช้ข้อมูลฝึก 80% และ median imputation ก่อนแบ่งข้อมูล ซึ่งเป็นข้อจำกัดด้าน leakage การเปรียบเทียบหลักด้านล่างใช้ผลรันใหม่ในสภาพแวดล้อมเดียวกัน ผล baseline ที่ผู้จัดทำเคยรายงานแยกเก็บใน historical_results.json และไม่ใช้แทน predictions ในการทดสอบนัยสำคัญ")
    p("ประเมิน paired bootstrap และ paired prediction-randomization อย่างละ 2,000 รอบบน test rows เดียวกัน โดยใช้ F1 gain และ RMSE reduction พร้อมปรับ p-value แบบ Bonferroni สำหรับสองงาน ผลนี้มีเงื่อนไขบนโมเดลที่ฝึกแล้ว สมมติให้รายการเป็นอิสระ และไม่ได้ปรับ survey weights หรือ clusters อีกทั้ง test set เคยถูกตรวจในรอบก่อน จึงเป็น follow-up benchmark ไม่ใช่ external validation")
    doc.add_page_break()
    h("4 Findings and Discussion")
    table=doc.add_table(rows=1,cols=3)
    for cell,text in zip(table.rows[0].cells,["ตัวชี้วัด","Baseline รันใหม่","โมเดลปรับปรุง"]):cell.text=text
    for task,d,keys in [("Classification",c,["accuracy","precision","recall","f1"]),("Regression",r,["rmse","mae","r2"])]:
        for k in keys:
            row=table.add_row().cells
            for cell,text in zip(row,[k.upper(),format(d["baseline"][k],",.4f"),format(d["improved"][k],",.4f")]):cell.text=text
    for i,row in enumerate(table.rows):
        for cell in row.cells:
            tc=cell._tc.get_or_add_tcPr();b=OxmlElement("w:tcBorders")
            for edge in ["top","left","bottom","right"]:
                e=OxmlElement("w:"+edge);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"D9D9D9");b.append(e)
            tc.append(b)
            if i==0:
                sh=OxmlElement("w:shd");sh.set(qn("w:fill"),"E9EEF2");tc.append(sh)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(2)
                for run in par.runs:run.font.size=Pt(10)
    cm=c["confusion_matrix"]
    p(f'Classification ตรวจพบกลุ่มยากจน {cm[1][1]} จาก {sum(cm[1])} ราย และ false positives {cm[0][1]} ราย ค่า threshold ที่เลือกคือ {c["threshold"]:.4f} ส่วน regression มี MAE {r["improved"]["mae"]:,.2f} ผลตัวเลขทุกค่ามาจากการรันนี้ การเพิ่มบาง metric ไม่ได้ยืนยันว่าโมเดลดีที่สุดในทุกเป้าหมาย')
    for name,d in [("F1 gain",c),("RMSE reduction",r)]:
        st=d["statistics"];lo,hi=st["gain_ci95"]
        verdict="พบหลักฐานว่าดีขึ้นภายใต้สมมติฐานที่ใช้" if st["observed_gain"]>0 and st["bonferroni_p"]<.05 else "ยังไม่ยืนยันว่าดีขึ้นอย่างมีนัยสำคัญ"
        p(f'{name}: ผลต่าง {st["observed_gain"]:.4f}, ช่วงความเชื่อมั่น 95% [{lo:.4f}, {hi:.4f}], adjusted p={st["bonferroni_p"]:.4f}; {verdict} ช่วงความเชื่อมั่นเป็นรายงาน ไม่ใช่ simultaneous intervals')
    p("การนำไปใช้จริงควรเริ่มจากเครื่องมือช่วยวิเคราะห์และให้เจ้าหน้าที่ตรวจสอบ ไม่ใช้ตัดสิทธิอัตโนมัติ หากต้องการ Recall สูงขึ้น ให้ทดลองลด threshold จากค่าที่เลือก โดยค้นหาช่วง 0.01–0.99 ทีละ 0.01 บน calibration data แล้วเลือกตามเป้าหมาย Recall และภาระการตรวจสอบ false positives ไฟล์ calibration_thresholds.csv แสดงผลที่คำนวณจริงบน calibration เท่านั้น ห้ามนำ test มาปรับเกณฑ์ และคะแนนจาก weighted loss ไม่ควรถูกตีความเป็นความน่าจะเป็นที่ผ่านการสอบเทียบแล้ว")
    p("Regression ใช้เป็นค่าประมาณประกอบการวิเคราะห์ ยังไม่มีจำนวนคงที่ที่ควรบวกหรือลบจาก prediction ทุกคน หากข้อมูลใหม่มี bias ให้ประมาณค่าเฉลี่ย y−prediction บน calibration ใหม่ แล้วทดสอบการชดเชยบนข้อมูลแยก ไม่ใช้ MAE/RMSE เป็นค่าชดเชย ก่อนใช้งานต้องตรวจความคลาดเคลื่อนรายกลุ่มและข้อมูลต่างพื้นที่หรือเวลา ไม่รับประกันว่าจะได้ผลเท่าชุดทดลอง")
    p("แหล่งข้อมูล: SES2023.csv และไฟล์ FFN ต้นฉบับใน repository Saelee-Kittipop/Deep-Learnning-LAB-Coding; รายละเอียดการรัน ผลดิบ และข้อจำกัดอยู่ใน README_TH.md")
    for border in list(doc.styles.element.iter(qn("w:pBdr")))+list(doc.element.iter(qn("w:pBdr"))):border.getparent().remove(border)
    doc.save(OUT/"LAB_FFN_Report_TH.docx")
    print(OUT/"LAB_FFN_Report_TH.docx")
if __name__=="__main__":create()
